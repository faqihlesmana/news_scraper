"""
app.py – Flask Dashboard untuk Pohuwato News Scraper
"""
import threading
import sys
import os

sys.path.insert(0, os.path.dirname(__file__))

from flask import Flask, render_template, jsonify, request
from config import FLASK_SECRET_KEY, FLASK_PORT, FLASK_DEBUG, SOURCES
from database.db import (
    init_db, get_db, save_article_feedback, get_feedback_by_article_ids,
    get_quarterly_entities, get_entity_references
)

_BASE_DIR = os.path.dirname(os.path.abspath(__file__))
app = Flask(
    __name__,
    template_folder=os.path.join(_BASE_DIR, "templates"),
    static_folder=os.path.join(_BASE_DIR, "static")
)
app.secret_key = FLASK_SECRET_KEY

_scrape_running = False
_backfill_running = False


# ─────────────────────────────────────────────────────────────────────────────
#  Helpers
# ─────────────────────────────────────────────────────────────────────────────
def db_query(sql, params=()):
    with get_db() as conn:
        rows = conn.execute(sql, params).fetchall()
        return [dict(r) for r in rows]


def db_one(sql, params=()):
    with get_db() as conn:
        row = conn.execute(sql, params).fetchone()
        return dict(row) if row else None


# ─────────────────────────────────────────────────────────────────────────────
#  Routes
# ─────────────────────────────────────────────────────────────────────────────
@app.route("/")
def index():
    return render_template("index.html", sources=SOURCES)


@app.route("/favicon.ico")
def favicon():
    return "", 204


@app.route("/triwulan")
def page_triwulan():
    """Halaman Dashboard Analisis Event & Lembaga Triwulanan."""
    return render_template("triwulan.html")


@app.route("/api/triwulan/data")
def api_triwulan_data():
    """Mengambil data event dan perusahaan per triwulan dan tahun."""
    year = request.args.get("year", type=int)
    quarter = request.args.get("quarter", "").strip() or None
    data = get_quarterly_entities(year=year, quarter=quarter)
    return jsonify(data)


@app.route("/api/triwulan/references")
def api_triwulan_references():
    """Mengambil daftar berita sumber untuk Pop-Up modal."""
    entity_type = request.args.get("type", "").strip()  # 'event' atau 'company'
    entity_id = request.args.get("id", type=int)

    if not entity_id or entity_type not in ("event", "company"):
        return jsonify({"error": "Parameter type (event/company) dan id wajib diisi"}), 400

    data = get_entity_references(entity_type, entity_id)
    return jsonify(data)


@app.route("/api/triwulan/backfill", methods=["POST"])
def api_triwulan_backfill():
    """Menjalankan proses ekstraksi berita lama di background."""
    global _backfill_running
    if _backfill_running:
        return jsonify({"status": "already_running", "message": "Backfill sedang berjalan..."}), 409

    data = request.get_json(silent=True) or {}
    limit = data.get("limit")

    def _run():
        global _backfill_running
        _backfill_running = True
        try:
            from backfill_entities import run_backfill
            run_backfill(limit=limit)
        finally:
            _backfill_running = False

    thread = threading.Thread(target=_run, daemon=True)
    thread.start()
    return jsonify({"status": "started", "message": "Proses ekstraksi berita lama dimulai di background..."})


@app.route("/api/triwulan/backfill/status")
def api_triwulan_backfill_status():
    return jsonify({"running": _backfill_running})


@app.route("/api/stats")
def api_stats():
    total        = db_one("SELECT COUNT(*) AS c FROM news")["c"]
    today        = db_one("SELECT COUNT(*) AS c FROM news WHERE DATE(scraped_at) = CURRENT_DATE")["c"]
    per_source   = db_query("SELECT source_media AS name, COUNT(*) AS count FROM news GROUP BY source_media ORDER BY count DESC")
    per_category = db_query("SELECT COALESCE(category, 'Belum Dikategorikan') AS name, COUNT(*) AS count FROM news GROUP BY name ORDER BY count DESC")
    recent_logs  = db_query("SELECT * FROM scrape_logs ORDER BY run_at DESC LIMIT 10")

    return jsonify({
        "total":        total,
        "today":        today,
        "per_source":   per_source,
        "per_category": per_category,
        "recent_logs":  recent_logs,
    })


@app.route("/api/news")
def api_news():
    page     = int(request.args.get("page", 1))
    per_page = int(request.args.get("per_page", 20))
    search   = request.args.get("search", "").strip()
    source   = request.args.get("source", "").strip()
    category   = request.args.get("category", "").strip()
    exp_cat    = request.args.get("exp_category", "").strip()
    start_date = request.args.get("start_date", "").strip()
    end_date   = request.args.get("end_date", "").strip()
    client_id  = request.headers.get("X-Client-ID") or request.args.get("client_id", "").strip() or request.remote_addr

    hide_irrelevant = request.args.get("hide_irrelevant", "").strip() == "1"

    offset = (page - 1) * per_page
    conditions = []
    params: list = []

    if search:
        conditions.append("(title LIKE ? OR content LIKE ?)")
        params.extend([f"%{search}%", f"%{search}%"])
    if source:
        conditions.append("source_slug = ?")
        params.append(source)
    if category:
        conditions.append("category = ?")
        params.append(category)
    if exp_cat:
        conditions.append("exp_category = ?")
        params.append(exp_cat)
    if start_date:
        conditions.append("SUBSTR(COALESCE(published_at, CAST(scraped_at AS TEXT)), 1, 10) >= ?")
        params.append(start_date)
    if end_date:
        conditions.append("SUBSTR(COALESCE(published_at, CAST(scraped_at AS TEXT)), 1, 10) <= ?")
        params.append(end_date)
    if hide_irrelevant:
        # Exclude articles where BOTH LU and PE are irrelevant/null
        conditions.append(
            "NOT ("
            "  (category IS NULL OR category = 'Tidak Relevan')"
            "  AND"
            "  (exp_category IS NULL OR exp_category = 'Tidak Relevan')"
            ")"
        )

    where = ("WHERE " + " AND ".join(conditions)) if conditions else ""

    count_sql = f"SELECT COUNT(*) AS c FROM news {where}"
    total = db_one(count_sql, params)["c"]

    news_sql = (
        f"SELECT id, url, source_media, title, author, published_at, "
        f"tags, image_url, category, subcategory, exp_category, exp_subcategory, "
        f"confidence, summary, scraped_at "
        f"FROM news {where} "
        f"ORDER BY published_at DESC NULLS LAST, scraped_at DESC "
        f"LIMIT ? OFFSET ?"
    )
    articles = db_query(news_sql, params + [per_page, offset])

    # Attach feedback data
    article_ids = [a["id"] for a in articles]
    fb_map = get_feedback_by_article_ids(article_ids, client_id)
    for a in articles:
        fb_info = fb_map.get(a["id"], {"sip": 0, "gasip": 0, "user_flag": None, "user_count": 0})
        a["feedback"] = fb_info

    return jsonify({
        "articles": articles,
        "total":    total,
        "page":     page,
        "per_page": per_page,
        "pages":    (total + per_page - 1) // per_page,
    })


@app.route("/api/news/<int:article_id>/feedback", methods=["POST"])
def api_save_feedback(article_id):
    data = request.get_json() or {}
    client_id = request.headers.get("X-Client-ID") or data.get("client_id") or request.remote_addr
    user_ip = request.remote_addr

    flag = data.get("flag")
    if flag is None or flag not in (0, 1):
        return jsonify({"success": False, "error": "Flag feedback tidak valid (harus 0 = Sip atau 1 = Ga Sip)."}), 400

    rec_cat = data.get("rec_category")
    rec_subcat = data.get("rec_subcategory")
    rec_exp_cat = data.get("rec_exp_category")
    rec_exp_subcat = data.get("rec_exp_subcategory")

    res = save_article_feedback(
        article_id=article_id,
        user_identifier=client_id,
        user_ip=user_ip,
        flag=flag,
        rec_cat=rec_cat,
        rec_subcat=rec_subcat,
        rec_exp_cat=rec_exp_cat,
        rec_exp_subcat=rec_exp_subcat,
        max_limit=3
    )

    if not res.get("success"):
        return jsonify(res), 429  # 429 Too Many Requests

    return jsonify(res)


@app.route("/api/news/<int:article_id>")
def api_article_detail(article_id):
    article = db_one("SELECT * FROM news WHERE id = ?", (article_id,))
    if not article:
        return jsonify({"error": "Not found"}), 404
    return jsonify(article)


@app.route("/api/sources")
def api_sources():
    return jsonify(SOURCES)


@app.route("/api/categories")
def api_categories():
    cats = db_query(
        "SELECT DISTINCT category FROM news WHERE category IS NOT NULL "
        "ORDER BY category"
    )
    return jsonify([r["category"] for r in cats])


@app.route("/api/ollama/status")
def api_ollama_status():
    """Cek apakah Ollama server aktif atau berjalan di Cloud Mode (Vercel)."""
    if os.environ.get("VERCEL"):
        return jsonify({
            "running": True,
            "cloud_mode": True,
            "url": "Supabase Cloud",
            "configured_model": "Cloud (Gemini API)",
            "model_ready": True,
            "available_models": ["gemini-flash-lite"],
        })

    from services.ollama_processor import is_ollama_running, get_available_models
    from config import OLLAMA_MODEL, OLLAMA_BASE_URL
    running = is_ollama_running()
    models  = get_available_models() if running else []
    return jsonify({
        "running":       running,
        "cloud_mode":     False,
        "url":           OLLAMA_BASE_URL,
        "configured_model": OLLAMA_MODEL,
        "model_ready":   any(OLLAMA_MODEL in m for m in models),
        "available_models": models,
    })


@app.route("/api/scrape", methods=["POST"])
def api_scrape():
    global _scrape_running
    if _scrape_running:
        return jsonify({"status": "already_running", "message": "Scraping sedang berjalan..."}), 409

    data         = request.get_json(silent=True) or {}
    source_slugs = data.get("sources")          # None = semua sumber
    use_ai       = data.get("use_ai", True)      # Ollama aktif by default
    use_gemini   = data.get("gemini", False)     # Override ke Gemini (fallback)
    max_pages    = int(data.get("pages", 3))

    def _run():
        global _scrape_running
        _scrape_running = True
        try:
            from services.scraper_engine import scrape_all
            scrape_all(max_pages=max_pages, use_ai=use_ai,
                       use_gemini=use_gemini, source_slugs=source_slugs)
        finally:
            _scrape_running = False

    thread = threading.Thread(target=_run, daemon=True)
    thread.start()
    return jsonify({"status": "started", "message": "Scraping dimulai di background..."})


@app.route("/api/scrape/status")
def api_scrape_status():
    return jsonify({"running": _scrape_running})


@app.route("/api/export/excel")
def api_export_excel():
    import io
    from datetime import datetime
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from flask import send_file

    search   = request.args.get("search", "").strip()
    source   = request.args.get("source", "").strip()
    category   = request.args.get("category", "").strip()
    exp_cat    = request.args.get("exp_category", "").strip()
    start_date = request.args.get("start_date", "").strip()
    end_date   = request.args.get("end_date", "").strip()

    conditions = []
    params: list = []

    if search:
        conditions.append("(title LIKE ? OR content LIKE ?)")
        params.extend([f"%{search}%", f"%{search}%"])
    if source:
        conditions.append("source_slug = ?")
        params.append(source)
    if category:
        conditions.append("category = ?")
        params.append(category)
    if exp_cat:
        conditions.append("exp_category = ?")
        params.append(exp_cat)
    if start_date:
        conditions.append("DATE(COALESCE(published_at, scraped_at)) >= DATE(?)")
        params.append(start_date)
    if end_date:
        conditions.append("DATE(COALESCE(published_at, scraped_at)) <= DATE(?)")
        params.append(end_date)

    where = ("WHERE " + " AND ".join(conditions)) if conditions else ""

    news_sql = (
        f"SELECT id, url, source_media, title, author, published_at, "
        f"content, summary, tags, image_url, category, subcategory, exp_category, exp_subcategory, "
        f"confidence, scraped_at "
        f"FROM news {where} "
        f"ORDER BY published_at DESC NULLS LAST, scraped_at DESC"
    )
    articles = db_query(news_sql, params)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Berita Pohuwato PDRB"
    ws.views.sheetView[0].showGridLines = True

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    data_font = Font(name="Calibri", size=10)
    data_align_left = Alignment(horizontal="left", vertical="top", wrap_text=True)
    data_align_center = Alignment(horizontal="center", vertical="top")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    headers = [
        "No", "Judul Berita", "Sumber Media", "Penulis", "Tanggal Publikasi",
        "PDRB Lapangan Usaha (Sektor)", "Sub-Sektor Lapangan Usaha",
        "PDRB Pengeluaran (Sektor)", "Sub-Sektor Pengeluaran",
        "Akurasi AI", "Ringkasan AI", "URL Berita", "Tanggal Scrape"
    ]
    ws.append(headers)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    ws.row_dimensions[1].height = 28

    for i, a in enumerate(articles, 1):
        conf_val = a.get("confidence")
        conf_str = f"{int(conf_val*100)}%" if conf_val is not None else "—"

        row_data = [
            i,
            a.get("title") or "—",
            a.get("source_media") or "—",
            a.get("author") or "—",
            a.get("published_at") or "—",
            a.get("category") or "Tidak Relevan",
            a.get("subcategory") or "Tidak Relevan",
            a.get("exp_category") or "Tidak Relevan",
            a.get("exp_subcategory") or "Tidak Relevan",
            conf_str,
            a.get("summary") or "",
            a.get("url") or "",
            a.get("scraped_at") or "",
        ]
        ws.append(row_data)

        row_num = i + 1
        ws.row_dimensions[row_num].height = 36

        for col_idx in range(1, len(row_data) + 1):
            c = ws.cell(row=row_num, column=col_idx)
            c.font = data_font
            c.border = thin_border
            if col_idx in (1, 5, 10, 13):
                c.alignment = data_align_center
            else:
                c.alignment = data_align_left

    col_widths = {
        1: 6, 2: 45, 3: 18, 4: 16, 5: 18,
        6: 25, 7: 25, 8: 22, 9: 25, 10: 12,
        11: 45, 12: 35, 13: 18
    }
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Berita_Pohuwato_PDRB_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename
    )


# ─────────────────────────────────────────────────────────────────────────────
#  Main
# ─────────────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    init_db()
    app.run(host="0.0.0.0", port=FLASK_PORT, debug=FLASK_DEBUG)
