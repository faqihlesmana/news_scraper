"""
services/gemini_categorizer.py

Alur Kategorisasi Dual PDRB Gemini:
  1. Setiap artikel di-clean & di-vectorize → keyword compact string
  2. Artikel dikirim ke Gemini (batch atau single API call)
  3. Gemini mengklasifikasikan ke 2 Perspektif PDRB:
     a. PDRB Lapangan Usaha (Kamus_Kategori_Berita_PDRB.xlsx)
     b. PDRB Pengeluaran (Kategorisasi Pengeluaran.xlsx)
  4. Mengembalikan: category, subcategory, exp_category, exp_subcategory, confidence, reasoning, is_relevant
"""
import json
import logging
import os
import re
from typing import Optional

from config import GEMINI_API_KEY, GEMINI_MODEL, EXCEL_CATEGORIES_PATH, EXCEL_PENGELUARAN_PATH
from services.text_vectorizer import vectorize_article

logger = logging.getLogger(__name__)

# Pemetaan Sektor PDRB Lapangan Usaha (BPS)
PDRB_LU_SECTOR_MAP = {
    'A':    'Pertanian, Kehutanan & Perikanan',
    'B':    'Pertambangan & Penggalian',
    'C':    'Industri Pengolahan',
    'D':    'Pengadaan Listrik & Gas',
    'E':    'Pengadaan Air, Pengelolaan Sampah & Daur Ulang',
    'F':    'Konstruksi',
    'G':    'Perdagangan Besar, Eceran & Reparasi',
    'H':    'Transportasi & Pergudangan',
    'I':    'Penyediaan Akomodasi & Makan Minum',
    'J':    'Informasi & Komunikasi',
    'K':    'Jasa Keuangan & Asuransi',
    'L':    'Real Estate',
    'MN':   'Jasa Perusahaan',
    'M':    'Jasa Perusahaan',
    'N':    'Jasa Perusahaan',
    'O':    'Administrasi Pemerintahan, Pertahanan & Sosial Wajib',
    'P':    'Jasa Pendidikan',
    'Q':    'Jasa Kesehatan & Kegiatan Sosial',
    'RSTU': 'Jasa Lainnya',
    'R':    'Jasa Lainnya',
    'S':    'Jasa Lainnya',
    'T':    'Jasa Lainnya',
    'U':    'Jasa Lainnya',
}

# 1. PDRB Lapangan Usaha Categories
LU_CATEGORIES: list[tuple[str, list[str]]] = []
_LU_CAT_NAMES: list[str] = []
_LU_SUBS_BY_CAT: dict[str, list[str]] = {}

_EXP_METADATA: dict[str, dict] = {}

def load_categories_from_excel() -> None:
    """
    Load both PDRB Lapangan Usaha and PDRB Pengeluaran categories from Excel files.
    """
    global LU_CATEGORIES, _LU_CAT_NAMES, _LU_SUBS_BY_CAT
    global EXP_CATEGORIES, _EXP_CAT_NAMES, _EXP_SUBS_BY_CAT, _EXP_METADATA

    import openpyxl

    # --- Load Lapangan Usaha ---
    if os.path.exists(EXCEL_CATEGORIES_PATH):
        try:
            wb = openpyxl.load_workbook(EXCEL_CATEGORIES_PATH, data_only=True)
            ws = wb.active
            lu_dict: dict[str, list[str]] = {}

            for i in range(2, ws.max_row + 1):
                kode = str(ws.cell(row=i, column=1).value or '').strip()
                name = str(ws.cell(row=i, column=2).value or '').strip()
                if not kode or not name:
                    continue
                prefix = ''.join([c for c in kode if c.isalpha()])
                main_cat = PDRB_LU_SECTOR_MAP.get(
                    prefix,
                    PDRB_LU_SECTOR_MAP.get(prefix[0] if prefix else '', 'Tidak Relevan')
                )
                sub_display = f"[{kode}] {name}"
                if sub_display not in lu_dict.setdefault(main_cat, []):
                    lu_dict[main_cat].append(sub_display)

            lu_dict["Tidak Relevan"] = ["Tidak Relevan"]
            LU_CATEGORIES = [(c, subs) for c, subs in lu_dict.items()]
            _LU_CAT_NAMES = [c for c, _ in LU_CATEGORIES]
            _LU_SUBS_BY_CAT = {c: subs for c, subs in LU_CATEGORIES}
            logger.info(f"[Gemini] Loaded {len(LU_CATEGORIES)} PDRB Lapangan Usaha categories")
        except Exception as e:
            logger.error(f"[Gemini] Error loading Lapangan Usaha Excel: {e}")
            _load_default_lu()
    else:
        _load_default_lu()

    # --- Load Pengeluaran (with Deskripsi & Keyword) ---
    if os.path.exists(EXCEL_PENGELUARAN_PATH):
        try:
            wb = openpyxl.load_workbook(EXCEL_PENGELUARAN_PATH, data_only=True)
            ws = wb.active
            exp_dict: dict[str, list[str]] = {}
            _EXP_METADATA.clear()

            for i in range(2, ws.max_row + 1):
                cat  = str(ws.cell(row=i, column=2).value or '').strip()
                sub  = str(ws.cell(row=i, column=3).value or '').strip()
                desc = str(ws.cell(row=i, column=5).value or '').strip()
                kw   = str(ws.cell(row=i, column=6).value or '').strip()

                if cat and sub:
                    if sub not in exp_dict.setdefault(cat, []):
                        exp_dict[cat].append(sub)
                    _EXP_METADATA[sub] = {"desc": desc, "keywords": kw}

            exp_dict["Tidak Relevan"] = ["Tidak Relevan"]
            EXP_CATEGORIES = [(c, subs) for c, subs in exp_dict.items()]
            _EXP_CAT_NAMES = [c for c, _ in EXP_CATEGORIES]
            _EXP_SUBS_BY_CAT = {c: subs for c, subs in EXP_CATEGORIES}
            logger.info(f"[Gemini] Loaded {len(EXP_CATEGORIES)} PDRB Pengeluaran categories (with descriptions & keywords)")
        except Exception as e:
            logger.error(f"[Gemini] Error loading Pengeluaran Excel: {e}")
            _load_default_exp()
    else:
        _load_default_exp()


def _load_default_lu():
    global LU_CATEGORIES, _LU_CAT_NAMES, _LU_SUBS_BY_CAT
    LU_CATEGORIES = [
        ("Pertanian, Kehutanan & Perikanan", ["Tanaman Pangan", "Peternakan", "Perikanan", "Kehutanan"]),
        ("Pertambangan & Penggalian", ["Pertambangan Bijih Logam", "Penggalian Lainnya"]),
        ("Industri Pengolahan", ["Industri Makanan dan Minuman", "Pengolahan Hasil Bumi"]),
        ("Konstruksi", ["Konstruksi Jalan & Bangunan"]),
        ("Perdagangan Besar, Eceran & Reparasi", ["Perdagangan Eceran", "Pasar"]),
        ("Transportasi & Pergudangan", ["Angkutan Darat", "Angkutan Laut"]),
        ("Informasi & Komunikasi", ["Telekomunikasi", "Internet"]),
        ("Administrasi Pemerintahan, Pertahanan & Sosial Wajib", ["Kebijakan Pemda", "Pelayanan Publik", "Ketertiban"]),
        ("Jasa Pendidikan", ["Sekolah", "Perguruan Tinggi"]),
        ("Jasa Kesehatan & Kegiatan Sosial", ["Rumah Sakit", "Layanan Kesehatan"]),
        ("Tidak Relevan", ["Tidak Relevan"]),
    ]
    _LU_CAT_NAMES   = [c for c, _ in LU_CATEGORIES]
    _LU_SUBS_BY_CAT = {c: subs for c, subs in LU_CATEGORIES}


def _load_default_exp():
    global EXP_CATEGORIES, _EXP_CAT_NAMES, _EXP_SUBS_BY_CAT
    EXP_CATEGORIES = [
        ("Konsumsi Rumah Tangga", ["PKRT Makanan dan Minuman Non Beralkohol", "PKRT Pakaian", "PKRT Perumahan"]),
        ("Konsumsi LNPRT", ["Pengeluaran Konsumsi LNPRT"]),
        ("Konsumsi Pemerintah", ["PKP Konsumsi Kolektif", "PKP Konsumsi Individu"]),
        ("PMTB", ["PMTB Bangunan", "PMTB Non-Bangunan"]),
        ("Perubahan Inventori", ["Perubahan Inventori"]),
        ("Ekspor", ["Ekspor LN Barang", "Ekspor LN Jasa", "Ekspor Antar Daerah"]),
        ("Impor", ["Impor LN Barang", "Impor LN Jasa", "Impor Antar Daerah"]),
        ("Tidak Relevan", ["Tidak Relevan"]),
    ]
    _EXP_CAT_NAMES   = [c for c, _ in EXP_CATEGORIES]
    _EXP_SUBS_BY_CAT = {c: subs for c, subs in EXP_CATEGORIES}


# Auto-load on import
load_categories_from_excel()


# ── Lazy API client ───────────────────────────────────────────────────────────
_client = None

def _get_client():
    global _client
    if _client is not None:
        return _client
    try:
        import google.genai as genai
        if not GEMINI_API_KEY:
            logger.warning("[Gemini] GEMINI_API_KEY not set in .env")
            return None
        _client = genai.Client(api_key=GEMINI_API_KEY)
        logger.info("[Gemini] Client initialized")
        return _client
    except ImportError:
        logger.warning("[Gemini] google-genai not installed.")
        return None


# ── System instruction ────────────────────────────────────────────────────────
def _build_system_instruction() -> str:
    lu_lines = [f"- {cat}: {', '.join(subs[:6])}" for cat, subs in LU_CATEGORIES]
    
    exp_lines = []
    for cat, subs in EXP_CATEGORIES:
        if cat == "Tidak Relevan":
            exp_lines.append(f"- {cat}: Tidak Relevan")
            continue
        sub_details = []
        for s in subs:
            meta = _EXP_METADATA.get(s, {})
            desc = meta.get("desc", "")
            kw = meta.get("keywords", "")
            info = f"  * {s}"
            if desc:
                info += f" -> {desc}"
            if kw:
                info += f" (Kata Kunci: {kw})"
            sub_details.append(info)
        exp_lines.append(f"- {cat}:\n" + "\n".join(sub_details))

    return (
        "Kamu adalah sistem klasifikasi berita sektor PDRB Kabupaten Pohuwato, Indonesia.\n"
        "Tugasmu mengklasifikasikan berita ke dalam 2 PERSPEKTIF PDRB:\n\n"
        "1. DUKUNGAN PDRB LAPANGAN USAHA:\n" + "\n".join(lu_lines) + "\n\n"
        "2. DUKUNGAN PDRB PENGELUARAN (dengan deskripsi & kata kunci):\n" + "\n".join(exp_lines) + "\n\n"
        "ATURAN RELEVANSI — SKALA DAMPAK EKONOMI:\n"
        "Yang dicari adalah berita yang BERDAMPAK SIGNIFIKAN terhadap perekonomian daerah.\n\n"
        "MASUK KATEGORI PDRB (r=true) jika:\n"
        "- Menyebut ANGKA PRODUKSI, NILAI EKONOMI, atau VOLUME yang jelas\n"
        "- PROGRAM/KEBIJAKAN PEMERINTAH berskala besar (bantuan sosial massal, APBD, pembangunan infrastruktur, subsidi)\n"
        "- BENCANA/KECELAKAAN BERSKALA BESAR yang berdampak luas pada ekonomi (banjir merusak ratusan hektar sawah, longsor menutup jalan utama)\n"
        "- KORUPSI/KASUS HUKUM BERSKALA BESAR yang melibatkan anggaran publik signifikan (korupsi miliaran)\n"
        "- AKTIVITAS EKONOMI yang melibatkan BANYAK ORANG/SEKTOR (panen raya, ekspor komoditas, pasar tradisional)\n"
        "- INVESTASI, PEMBUKAAN/PENUTUPAN USAHA yang berdampak pada lapangan kerja\n\n"
        "TIDAK RELEVAN / 'Tidak Relevan' (r=false) jika:\n"
        "- Kecelakaan/kematian INDIVIDU tanpa dampak ekonomi luas (1 orang terjatuh, kecelakaan tunggal)\n"
        "- Kriminal SKALA KECIL (pencurian individu, perkelahian, penipuan perorangan)\n"
        "- Berita seremonial/protokoler tanpa substansi ekonomi\n"
        "- Hiburan, gosip, olahraga, agama murni tanpa kaitan ekonomi\n"
        "- Berita sosial individu (pernikahan, kelahiran, kematian biasa)\n\n"
        "CONTOH:\n"
        "\"Penambang di DAM Tewas Terjatuh\" → r:false, Tidak Relevan (kecelakaan individu)\n"
        "\"Banjir Rendam 500 Ha Sawah, Rugi Miliaran\" → r:true, Pertanian (bencana besar, dampak ekonomi)\n"
        "\"Pemuda Ditangkap Mencuri Motor\" → r:false, Tidak Relevan (kriminal kecil)\n"
        "\"KPK Usut Korupsi Dana Desa Rp 3,2 M\" → r:true, Administrasi Pemerintahan (korupsi besar)\n"
        "\"Pemerintah Salurkan 10.000 Paket Sembako\" → r:true, Konsumsi Pemerintah (bantuan massal)\n"
        "\"Satu Unit Rumah Terbakar\" → r:false, Tidak Relevan (insiden kecil individual)\n\n"
        "ATURAN OUTPUT:\n"
        "- 'c_lu' & 's_lu': Kategori & Sub-kategori Lapangan Usaha.\n"
        "- 'c_pe' & 's_pe': Kategori & Sub-kategori Pengeluaran.\n"
        "- 'r': true jika berita BERDAMPAK EKONOMI SIGNIFIKAN, false jika berita skala kecil/individu/non-ekonomi.\n"
        "- Jika r=false, set c_lu:'Tidak Relevan', s_lu:'Tidak Relevan', c_pe:'Tidak Relevan', s_pe:'Tidak Relevan'.\n"
        "- Output HANYA berupa JSON array valid tanpa spasi/newline ekstra:\n"
        "[{\"id\":\"<id>\",\"c_lu\":\"<kategori_lu>\",\"s_lu\":\"<subkategori_lu>\",\"c_pe\":\"<kategori_pe>\",\"s_pe\":\"<subkategori_pe>\",\"n\":<1-10>,\"r\":<true|false>}]"
    )


# ── Core batch API call ───────────────────────────────────────────────────────
def _batch_call_gemini(articles_payload: list[dict]) -> Optional[list[dict]]:
    client = _get_client()
    if not client:
        return None

    user_content = json.dumps(articles_payload, ensure_ascii=False, separators=(",", ":"))

    models_to_try = [GEMINI_MODEL, "gemini-flash-lite-latest", "gemini-flash-latest"]
    seen = set()
    models = [m for m in models_to_try if m and not (m in seen or seen.add(m))]

    import google.genai as genai
    for model_name in models:
        try:
            response = client.models.generate_content(
                model=model_name,
                contents=user_content,
                config=genai.types.GenerateContentConfig(
                    system_instruction=_build_system_instruction(),
                    temperature=0.0,
                    max_output_tokens=len(articles_payload) * 120 + 150,
                ),
            )
            raw = (response.text or "").strip()
            raw = re.sub(r"^```(?:json)?\s*", "", raw)
            raw = re.sub(r"\s*```$", "", raw)

            m = re.search(r"\[.*\]", raw, re.DOTALL)
            if m:
                raw = m.group(0)

            result = json.loads(raw)
            if isinstance(result, list):
                return result
            logger.error(f"[Gemini] Response from {model_name} is not a list: {raw[:200]}")
        except Exception as e:
            logger.warning(f"[Gemini] API error with model {model_name}: {e}")

    return None


# ── Validation helpers ────────────────────────────────────────────────────────
def _validate_lu(cat: str, sub: str) -> tuple[str, str]:
    if cat not in _LU_CAT_NAMES:
        cl = cat.lower()
        cat = next((v for v in _LU_CAT_NAMES if cl in v.lower() or v.lower() in cl), "Tidak Relevan")

    valid_subs = _LU_SUBS_BY_CAT.get(cat, ["Tidak Relevan"])
    if sub not in valid_subs:
        sl = sub.lower()
        sub = next((s for s in valid_subs if sl in s.lower() or s.lower() in sl), valid_subs[0])

    return cat, sub


def _validate_exp(cat: str, sub: str) -> tuple[str, str]:
    if cat not in _EXP_CAT_NAMES:
        cl = cat.lower()
        cat = next((v for v in _EXP_CAT_NAMES if cl in v.lower() or v.lower() in cl), "Tidak Relevan")

    valid_subs = _EXP_SUBS_BY_CAT.get(cat, ["Tidak Relevan"])
    if sub not in valid_subs:
        sl = sub.lower()
        sub = next((s for s in valid_subs if sl in s.lower() or s.lower() in sl), valid_subs[0])

    return cat, sub


# ── Public interface ──────────────────────────────────────────────────────────
def batch_categorize(articles: list[dict]) -> dict[str, dict]:
    if not articles:
        return {}

    payloads = []
    for i, art in enumerate(articles):
        art_id  = str(art.get("id") or art.get("url") or i)
        vec     = vectorize_article(art.get("title", ""), art.get("content", ""))
        payload = {
            "id":    art_id,
            "title": (art.get("title") or "")[:120],
            "vec":   vec,
        }
        payloads.append(payload)

    logger.info(f"[Gemini] Batch categorizing {len(payloads)} articles into Dual PDRB categories")
    raw_results = _batch_call_gemini(payloads)

    output: dict[str, dict] = {}
    if raw_results:
        for item in raw_results:
            try:
                art_id = str(item.get("id", ""))
                c_lu, s_lu = _validate_lu(item.get("c_lu", "Tidak Relevan"), item.get("s_lu", "Tidak Relevan"))
                c_pe, s_pe = _validate_exp(item.get("c_pe", "Tidak Relevan"), item.get("s_pe", "Tidak Relevan"))

                conf   = min(1.0, float(item.get("n", 5)) / 10.0)
                is_rel = bool(item.get("r", True)) if (c_lu != "Tidak Relevan" or c_pe != "Tidak Relevan") else False

                output[art_id] = {
                    "category":        c_lu,
                    "subcategory":     s_lu,
                    "exp_category":    c_pe,
                    "exp_subcategory": s_pe,
                    "confidence":      conf,
                    "reasoning":       "",
                    "is_relevant":     is_rel,
                }
            except Exception as e:
                logger.warning(f"[Gemini] Could not parse batch item {item}: {e}")

    for payload in payloads:
        art_id = payload["id"]
        if art_id not in output:
            output[art_id] = {
                "category":        "Tidak Relevan",
                "subcategory":     "Tidak Relevan",
                "exp_category":    "Tidak Relevan",
                "exp_subcategory": "Tidak Relevan",
                "confidence":      0.0,
                "reasoning":       "",
                "is_relevant":     False,
            }

    return output


def categorize(title: str, content: str) -> dict:
    results = batch_categorize([{"id": "_single", "title": title, "content": content}])
    return results.get("_single", {
        "category":        "Tidak Relevan",
        "subcategory":     "Tidak Relevan",
        "exp_category":    "Tidak Relevan",
        "exp_subcategory": "Tidak Relevan",
        "confidence":      0.0,
        "reasoning":       "",
        "is_relevant":     False,
    })
